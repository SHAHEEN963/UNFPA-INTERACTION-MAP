#!/usr/bin/env python3
"""Validate XLSX locally and generate only the public map projection.

No workbook values, file paths or exception strings are emitted in diagnostics.
There is deliberately no internal/unredacted export mode.
"""
from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile
import unicodedata
import zipfile

from openpyxl import load_workbook
from openpyxl.xml import DEFUSEDXML

ROOT = Path(__file__).resolve().parents[1]
SHEETS = ("Lattakia", "Tartous", "Hama", "Homs")
IGNORED_SHEETS = {"Sheet1", "Calc_Data", "Statistics"}
GOV_AR = {"Lattakia": "اللاذقية", "Tartous": "طرطوس", "Hama": "حماة", "Homs": "حمص"}
PROGRAMS = {"rh": "RH", "gbv": "GBV", "youth": "Youth"}
# This is a publication allowlist, not an exhaustive ontology of services.
PUBLIC_TYPES = {
    ("RH", "RH Static Clinic", "RH SC"),
    ("RH", "Breast Cancer Unit", "RH SC"),
    ("RH", "CEmONC", "RH SC"),
    ("Youth", "Youth Friendly Space", "YFS"),
}
HEADERS = {
    "implementing_partner": "*Final Implementing Partner (IP) الشريك المنفذ النهائي",
    "facility_code": "*Facility Code / رمز المنشأة",
    "main_program": "*Main Program",
    "facility_type": "*Main Facility Type / نوع المنشأة",
    "mode": "Mode of Facility",
    "governorate": "*Governorate",
    "district": "*District",
    "subdistrict": "*Sub District",
    "commune": "*Commune | Village | Town",
    "neighbourhood": "Neighbourhood",
    "site_name": "Site Name \nاسم الموقع (العنوان التفصيلي)",
    "lat": "Latitude",
    "lon": "Longtitude",  # spelling in the verified original workbook
    "brochure": "Detailed Information for the Brochure",
}
MAX_FILE = 20 * 1024 * 1024
MAX_UNPACKED = 100 * 1024 * 1024
MAX_ROWS = 20000
SWAP_FLAG = "Likely inverted latitude/longitude columns; map uses swapped values, field verification needed"
PREFIX = "/* PUBLIC edition: protected facility details are removed before serialization. */\nwindow.UNFPA_PUBLIC_MODE=true;\nwindow.UNFPA_FACILITIES="


class Invalid(Exception):
    """Contains only a fixed diagnostic code, never input values."""


def text(value):
    return "" if value is None else str(value).strip()


def normalized(value):
    return " ".join(unicodedata.normalize("NFKC", text(value)).casefold().split())


def header_key(value):
    return normalized(value).lstrip("* ")


def digest(value):
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def dumps(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False) + "\n"


def coordinate(value, axis):
    """Decimal degrees or strict DMS, with optional hemisphere and second mark."""
    if value is None or text(value) == "":
        raise Invalid("missing_coordinate")
    if isinstance(value, bool):
        raise Invalid("invalid_coordinate")
    s = unicodedata.normalize("NFKC", text(value)).upper().strip()
    try:
        result = float(s)
    except ValueError:
        match = re.fullmatch(
            r"([+-]?\d{1,3})\s*[°º]\s*(\d{1,2})\s*['′’]\s*"
            r"(\d{1,2}(?:\.\d+)?)\s*[\"″”]?\s*([NSEW])?", s
        )
        if not match:
            raise Invalid("invalid_coordinate") from None
        degree, minute, second, hemi = match.groups()
        d, m, sec = int(degree), int(minute), float(second)
        if m >= 60 or sec >= 60:
            raise Invalid("invalid_dms")
        if hemi and hemi not in ("NS" if axis == "lat" else "EW"):
            raise Invalid("invalid_hemisphere")
        if degree.startswith("-") and hemi in ("N", "E"):
            raise Invalid("conflicting_hemisphere")
        result = abs(d) + m / 60 + sec / 3600
        if degree.startswith("-") or hemi in ("S", "W"):
            result = -result
    limit = 90 if axis == "lat" else 180
    if not math.isfinite(result) or abs(result) > limit:
        raise Invalid("coordinate_out_of_range")
    return result


def in_extent(lat, lon):
    # Broad Syria sanity bounds, NOT a land mask or a field verification.
    return 32 <= lat <= 38 and 35.5 <= lon <= 43


def read_baseline(path):
    if path.stat().st_size > MAX_FILE:
        raise Invalid("baseline_too_large")
    raw = path.read_text(encoding="utf-8")
    match = re.fullmatch(
        r"\s*(?:/\*.*?\*/\s*)?window\.UNFPA_PUBLIC_MODE\s*=\s*true\s*;\s*"
        r"window\.UNFPA_FACILITIES\s*=\s*(\[.*\])\s*;\s*", raw, re.S
    )
    if not match:
        raise Invalid("invalid_public_baseline")
    try:
        records = json.loads(match[1])
    except (ValueError, TypeError):
        raise Invalid("invalid_public_baseline") from None
    seen = set()
    for r in records:
        if (not isinstance(r, dict) or not isinstance(r.get("id"), str)
                or r["id"] in seen or r.get("governorate_en") not in SHEETS
                or r.get("main_program") not in PROGRAMS.values()
                or not isinstance(r.get("facility_code"), str)):
            raise Invalid("invalid_public_baseline")
        seen.add(r["id"])
        if r.get("protected") or r["main_program"] == "GBV":
            if not r.get("protected") or r.get("lat") is not None or r.get("lon") is not None:
                raise Invalid("unsafe_public_baseline")
    if not records or len(records) > MAX_ROWS:
        raise Invalid("invalid_public_baseline")
    return raw, records


def read_policy(path):
    try:
        p = json.loads(path.read_text(encoding="utf-8"))
        if set(p) != {"version", "approved_public_codes"} or p["version"] != 1:
            raise ValueError
        codes = p["approved_public_codes"]
        if not isinstance(codes, list) or any(not isinstance(x, str) or not x.strip() for x in codes):
            raise ValueError
        result = {normalized(x) for x in codes}
        if len(result) != len(codes):
            raise ValueError
        return result
    except (ValueError, TypeError, KeyError):
        raise Invalid("invalid_public_policy") from None


def read_workbook(path):
    if path.suffix.lower() != ".xlsx":
        raise Invalid("xlsx_required")
    if path.stat().st_size > MAX_FILE:
        raise Invalid("workbook_too_large")
    if not DEFUSEDXML:
        raise Invalid("defusedxml_required")
    with zipfile.ZipFile(path) as z:
        entries = z.infolist()
        if len(entries) > 10000 or sum(x.file_size for x in entries) > MAX_UNPACKED:
            raise Invalid("workbook_expansion_limit")
        if any("vbaproject" in x.filename.lower() for x in entries):
            raise Invalid("macros_not_allowed")
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    records, errors = [], Counter()
    try:
        if not set(SHEETS).issubset(wb.sheetnames):
            raise Invalid("missing_required_sheet")
        if set(wb.sheetnames) - set(SHEETS) - IGNORED_SHEETS:
            raise Invalid("unexpected_sheet")
        aliases = {header_key(v): k for k, v in HEADERS.items()}
        aliases[header_key("Longitude")] = "lon"
        aliases[header_key("الرقم")] = "_number"
        for sheet in SHEETS:
            ws = wb[sheet]
            # Do not trust potentially stale Excel dimension metadata.
            ws.reset_dimensions()
            rows = ws.iter_rows()
            header = next(rows, ())
            if len(header) > 64:
                raise Invalid("column_limit")
            columns = {}
            for index, cell in enumerate(header):
                if not text(cell.value):
                    continue
                key = aliases.get(header_key(cell.value))
                if key is None:
                    errors["unknown_column"] += 1
                elif key in columns:
                    errors["duplicate_column"] += 1
                else:
                    columns[key] = index
            if not set(HEADERS).issubset(columns):
                errors["missing_required_column"] += 1
                continue
            if errors:
                continue
            sheet_count = 0
            for rownum, cells in enumerate(rows, start=2):
                if rownum > MAX_ROWS + 1:
                    raise Invalid("row_limit")
                if len(cells) > 64:
                    raise Invalid("column_limit")
                if not any(text(c.value) for c in cells):
                    continue
                # The original has an array formula in the unused serial-number
                # column. Never evaluate or import it; formulas in data are errors.
                data_cells = [c for i, c in enumerate(cells) if i != columns.get("_number")]
                if any(c.data_type == "f" for c in data_cells):
                    errors["formula_not_allowed"] += 1
                    continue
                if any(c.data_type == "e" for c in data_cells):
                    errors["excel_error"] += 1
                    continue
                if any(text(c.value) for i, c in enumerate(cells) if i not in columns.values()):
                    errors["unlabelled_data_column"] += 1
                    continue
                r = {k: cells[i].value if i < len(cells) else None for k, i in columns.items() if k != "_number"}
                r.update(source_sheet=sheet, source_row=rownum)
                records.append(r)
                sheet_count += 1
                if len(records) > MAX_ROWS:
                    raise Invalid("row_limit")
            if not sheet_count:
                errors["empty_required_sheet"] += 1
        return records, errors
    finally:
        wb.close()


def validate(records, previous):
    errors, warnings, output = Counter(), Counter(), []
    old = {normalized(r["facility_code"]): r for r in previous if not r.get("protected")}
    identities, duplicates = set(), set()
    for raw in records:
        r = {k: text(v) for k, v in raw.items() if k not in ("source_row", "lat", "lon")}
        r["source_row"] = raw["source_row"]
        for field in ("facility_code", "implementing_partner", "main_program", "facility_type", "governorate"):
            if not r[field]:
                errors["missing_" + field] += 1
        key = normalized(r["facility_code"])
        if key in identities:
            errors["duplicate_facility_code"] += 1
        identities.add(key)
        if key.startswith("protected-"):
            errors["reserved_facility_code"] += 1
        r["id"] = r["facility_code"]
        r["main_program"] = PROGRAMS.get(normalized(r["main_program"]))
        if r["main_program"] is None:
            errors["unknown_program"] += 1
        sheet = r["source_sheet"]
        if normalized(r["governorate"]) not in {normalized(sheet), normalized(GOV_AR[sheet]), normalized(sheet + "/" + GOV_AR[sheet])}:
            errors["governorate_sheet_mismatch"] += 1
        r["governorate_en"] = sheet
        r["governorate"] = sheet + "/" + GOV_AR[sheet]
        r["quality_flags"] = []
        r["lat_original"], r["lon_original"] = text(raw["lat"]), text(raw["lon"])
        try:
            lat, lon = coordinate(raw["lat"], "lat"), coordinate(raw["lon"], "lon")
            baseline = old.get(key, {})
            # Preserve ONLY an already flagged V4 correction with identical source values.
            # No heuristic swapping and no row-number exception.
            if SWAP_FLAG in baseline.get("quality_flags", []):
                same_source = (lat == coordinate(baseline["lat_original"], "lat")
                               and lon == coordinate(baseline["lon_original"], "lon"))
                if same_source and lat == baseline.get("lon") and lon == baseline.get("lat"):
                    lat, lon = lon, lat
                    r["quality_flags"].append(SWAP_FLAG)
                    warnings["existing_unverified_swap_retained"] += 1
            if not in_extent(lat, lon):
                raise Invalid("outside_syria_extent_review_source")
            r["lat"], r["lon"] = lat, lon
            if any("°" in text(raw[k]) for k in ("lat", "lon")):
                r["quality_flags"].append("DMS coordinates converted to decimal degrees")
                warnings["dms_converted"] += 1
        except Invalid as exc:
            errors[str(exc)] += 1
            continue
        # A different code on the same complete record is still a suspected duplicate.
        fingerprint = tuple(normalized(r[k]) for k in HEADERS if k != "facility_code")
        if fingerprint in duplicates:
            errors["duplicate_record_different_code"] += 1
        duplicates.add(fingerprint)
        if not r["site_name"]:
            r["quality_flags"].append("Detailed site name/address missing from source")
        output.append(r)
    return output, errors, warnings


def eligible(r, approved):
    triple = (r["main_program"], r["facility_type"], r["mode"])
    return normalized(r["facility_code"]) in approved and triple in PUBLIC_TYPES


def protected_record(gov, program, number):
    # Construct from constants. Never spread/copy the source record here.
    r = {k: "Withheld" for k in HEADERS}
    r.update(
        id=f"PROTECTED-{number:03d}", facility_code=f"Protected facility {number:02d}",
        main_program=program, facility_type="Protection services",
        governorate=gov + "/" + GOV_AR[gov], governorate_en=gov,
        site_name="Exact site information withheld on public map", lat=None, lon=None,
        lat_original="Withheld", lon_original="Withheld",
        brochure="Details are not published on the public website.",
        source_sheet=gov, source_row=None, quality_flags=[], protected=True,
    )
    return r


def public_projection(records, approved):
    public, groups = [], Counter()
    for r in records:
        if eligible(r, approved):
            # Explicit schema allowlist; unexpected workbook fields cannot be serialized.
            public.append({k: r[k] for k in (*HEADERS, "id", "governorate_en", "source_sheet",
                                           "source_row", "lat_original", "lon_original", "quality_flags")})
        else:
            groups[(r["governorate_en"], r["main_program"])] += 1
    public.sort(key=lambda r: normalized(r["id"]))
    co_locations = Counter((r["lat"], r["lon"]) for r in public)
    for r in public:
        n = co_locations[(r["lat"], r["lon"])]
        if n > 1:
            r["quality_flags"].append(f"Shares identical coordinates with {n - 1} other facility record(s); pins separated visually")
    number = 0
    # Synthetic IDs depend only on aggregate buckets, never source ID/order/row/hash.
    for (gov, program), count in sorted(groups.items()):
        for _ in range(count):
            number += 1
            public.append(protected_record(gov, program, number))
    return public


def counts(records):
    return {
        "total": len(records), "protected": sum(bool(r.get("protected")) for r in records),
        "mapped": sum(r.get("lat") is not None and r.get("lon") is not None for r in records),
        "by_governorate": {g: sum(r["governorate_en"] == g for r in records) for g in SHEETS},
        "by_program": {p: sum(r["main_program"] == p for r in records) for p in PROGRAMS.values()},
    }


def difference(previous, current):
    before = {normalized(r["facility_code"]): r for r in previous if not r.get("protected")}
    after = {normalized(r["facility_code"]): r for r in current if not r.get("protected")}
    changed, fields = 0, Counter()
    # Source row/order and formatting changes are not substantive facility updates.
    compare_fields = set(HEADERS) | {"governorate_en"}
    for key in before.keys() & after.keys():
        changed_fields = [f for f in compare_fields if before[key].get(f) != after[key].get(f)]
        changed += bool(changed_fields)
        fields.update(changed_fields)
    return {
        "before": counts(previous), "after": counts(current),
        "public_records_added_or_newly_published": len(after.keys() - before.keys()),
        "public_records_removed_or_newly_withheld": len(before.keys() - after.keys()),
        "public_records_changed": changed, "changed_public_field_counts": dict(sorted(fields.items())),
        "protected_comparison": "Aggregate counts only; individual protected additions, removals and edits cannot be inferred from a redacted baseline.",
    }


def serialize(records):
    return PREFIX + dumps(records).rstrip().replace("<", "\\u003c").replace("\u2028", "\\u2028").replace("\u2029", "\\u2029") + ";\n"


def atomic_write(path, content):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".import-", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as stream:
            stream.write(content)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)


def run(workbook, baseline, policy, review_dir, apply=False, approve=None):
    report_path = review_dir / "report.json"
    candidate_path = review_dir / "facilities-data.js"
    inputs = {p.resolve() for p in (workbook, baseline, policy)}
    if len(inputs) != 3 or any(p.resolve() in inputs for p in (report_path, candidate_path)):
        raise Invalid("input_output_path_collision")
    # Invalidate an earlier preview first: a failed run must not leave a stale candidate.
    candidate_path.unlink(missing_ok=True)
    report = {"status": "blocked", "errors": {}, "warnings": {}}
    try:
        raw_baseline, previous = read_baseline(baseline)
        approved = read_policy(policy)
        # Carry forward only already-public, non-protection approvals. No real
        # facility identifiers need to be shipped in the tool's default policy.
        approved |= {normalized(r["facility_code"]) for r in previous
                     if not r.get("protected")
                     and (r.get("main_program"), r.get("facility_type"), r.get("mode")) in PUBLIC_TYPES}
        raw_records, errors = read_workbook(workbook)
        records, validation_errors, warnings = validate(raw_records, previous)
        errors.update(validation_errors)
        report.update(errors=dict(sorted(errors.items())), warnings=dict(sorted(warnings.items())))
        if errors:
            atomic_write(report_path, dumps(report))
            return report, 2
        current = public_projection(records, approved)
        candidate = serialize(current)
        report["candidate_sha256"] = digest(candidate)
        report["diff"] = difference(previous, current)
        # Bind approval to both the exact baseline and the reviewed public candidate/policy.
        token = digest(raw_baseline + "\0" + candidate + "\0" + dumps(sorted(approved)))
        report["approval_token"] = token
        report["status"] = "preview"
        if report["diff"]["before"] != report["diff"]["after"]:
            report["warnings"]["facility_counts_changed_review_required"] = 1
        if apply and approve != token:
            report["status"] = "blocked"
            report["errors"]["approval_token_required_or_stale"] = 1
        # Persist the report before any candidate or published data write.
        atomic_write(report_path, dumps(report))
        if report["errors"]:
            return report, 2
        atomic_write(candidate_path, candidate)
        if apply:
            # Refuse to overwrite another process's intervening update.
            if baseline.read_text(encoding="utf-8") != raw_baseline:
                raise Invalid("baseline_changed_during_run")
            report["status"] = "applied"
            atomic_write(report_path, dumps(report))
            # The snapshot replacement is the final fallible write. A failed
            # report/candidate write must never leave a changed live snapshot.
            atomic_write(baseline, candidate)
        return report, 0
    except Invalid as exc:
        report["status"] = "blocked"
        report["errors"][str(exc)] = 1
        candidate_path.unlink(missing_ok=True)
        atomic_write(report_path, dumps(report))
        return report, 2


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--baseline", type=Path, default=ROOT / "assets/facilities-data.js")
    parser.add_argument("--policy", type=Path, default=ROOT / "tools/public-policy.json")
    parser.add_argument("--review-dir", type=Path, default=ROOT / "build/import-review")
    parser.add_argument("--apply", action="store_true", help="Atomically replace the baseline after reviewing the preview")
    parser.add_argument("--approve", help="Exact approval_token from the preview report")
    args = parser.parse_args(argv)
    try:
        report, code = run(args.workbook, args.baseline, args.policy, args.review_dir, args.apply, args.approve)
    except Exception:
        # Even malformed ZIP/XML and OS exceptions can contain source values/paths.
        # Never propagate traceback/exception text to CI, console or public reports.
        report, code = {"status": "blocked", "errors": {"input_or_output_failure": 1}}, 2
        try:
            # Do not overwrite input paths in an invalid path configuration.
            paths = {p.resolve() for p in (args.workbook, args.baseline, args.policy)}
            if (args.review_dir / "facilities-data.js").resolve() not in paths:
                (args.review_dir / "facilities-data.js").unlink(missing_ok=True)
            if (args.review_dir / "report.json").resolve() not in paths:
                atomic_write(args.review_dir / "report.json", dumps(report))
        except Exception:
            pass
    print(dumps(report), end="")
    return code


if __name__ == "__main__":
    sys.exit(main())
