"""All fixtures are synthetic and generated in temporary directories."""
from contextlib import redirect_stdout
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from tools import import_facilities as imp


class ImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.xlsx = self.root / 'source.xlsx'
        self.baseline = self.root / 'baseline.js'
        self.policy = self.root / 'policy.json'
        self.review = self.root / 'review'
        self.rows = {}
        self.old = []
        for i, gov in enumerate(imp.SHEETS):
            r = dict(implementing_partner='Test Partner', facility_code=f'TEST-{i}',
                     main_program='RH', facility_type='RH Static Clinic', mode='RH SC',
                     governorate=gov, district='District', subdistrict='Subdistrict',
                     commune='Town', neighbourhood='Area', site_name='Test site',
                     lat=round(35.2 + i / 100, 6), lon=round(36.5 + i / 100, 6), brochure='Test details')
            self.rows[gov] = [r]
            previous = dict(r, id=r['facility_code'], governorate_en=gov,
                            governorate=gov + '/' + imp.GOV_AR[gov],
                            source_sheet=gov, source_row=2, quality_flags=[],
                            lat_original=str(r['lat']), lon_original=str(r['lon']))
            self.old.append(previous)
        self.baseline.write_text(imp.serialize(self.old), encoding='utf-8')
        self.policy.write_text(imp.dumps({'version': 1, 'approved_public_codes': [r['id'] for r in self.old]}))
        self.write_workbook()

    def write_workbook(self):
        wb = Workbook()
        wb.remove(wb.active)
        for gov in imp.SHEETS:
            ws = wb.create_sheet(gov)
            ws.append(['الرقم', *imp.HEADERS.values()])
            for i, r in enumerate(self.rows[gov], start=1):
                ws.append([i, *(r[k] for k in imp.HEADERS)])
        wb.save(self.xlsx)
        wb.close()

    def mutate(self, callback):
        wb = load_workbook(self.xlsx)
        callback(wb)
        wb.save(self.xlsx)
        wb.close()

    def run_import(self, **kwargs):
        return imp.run(self.xlsx, self.baseline, self.policy, self.review, **kwargs)

    def assert_blocked(self, expected):
        before = self.baseline.read_bytes()
        report, code = self.run_import()
        self.assertEqual(code, 2)
        self.assertIn(expected, report['errors'])
        self.assertEqual(self.baseline.read_bytes(), before)
        self.assertFalse((self.review / 'facilities-data.js').exists())

    def test_preview_then_apply_requires_exact_token(self):
        before = self.baseline.read_bytes()
        report, code = self.run_import()
        self.assertEqual(code, 0)
        self.assertEqual(report['diff']['after']['total'], 4)
        self.assertEqual(self.baseline.read_bytes(), before)
        self.assertEqual(self.run_import(apply=True)[1], 2)
        report, _ = self.run_import()
        applied, code = self.run_import(apply=True, approve=report['approval_token'])
        self.assertEqual((applied['status'], code), ('applied', 0))
        self.assertEqual(self.baseline.read_bytes(), (self.review / 'facilities-data.js').read_bytes())

    def test_stale_approval_for_changed_source(self):
        report, _ = self.run_import()
        self.rows['Hama'][0]['site_name'] = 'Updated address'
        self.write_workbook()
        before = self.baseline.read_bytes()
        report, code = self.run_import(apply=True, approve=report['approval_token'])
        self.assertEqual(code, 2)
        self.assertIn('approval_token_required_or_stale', report['errors'])
        self.assertEqual(self.baseline.read_bytes(), before)

    def test_stale_approval_for_changed_baseline(self):
        report, _ = self.run_import()
        self.baseline.write_text(self.baseline.read_text() + '\n')
        self.assertEqual(self.run_import(apply=True, approve=report['approval_token'])[1], 2)

    def test_unknown_missing_and_duplicate_headers(self):
        for change, error in [(lambda s: setattr(s['M1'], 'value', 'Latitdue'), 'unknown_column'),
                              (lambda s: setattr(s['M1'], 'value', None), 'missing_required_column'),
                              (lambda s: setattr(s['N1'], 'value', 'Latitude'), 'duplicate_column')]:
            with self.subTest(error=error):
                self.write_workbook()
                self.mutate(lambda w: change(w['Lattakia']))
                self.assert_blocked(error)

    def test_header_whitespace_and_corrected_longitude_spelling(self):
        self.mutate(lambda w: setattr(w['Homs']['N1'], 'value', '  Longitude  '))
        self.assertEqual(self.run_import()[1], 0)

    def test_column_reorder_maps_by_name(self):
        def reorder(w):
            s = w['Lattakia']
            for row in s:
                row[1].value, row[2].value = row[2].value, row[1].value
        self.mutate(reorder)
        self.assertEqual(self.run_import()[0]['diff']['public_records_changed'], 0)

    def test_missing_sheet(self):
        self.mutate(lambda w: w.remove(w['Homs']))
        self.assert_blocked('missing_required_sheet')

    def test_unexpected_sheet(self):
        self.mutate(lambda w: w.create_sheet('Unreviewed'))
        self.assert_blocked('unexpected_sheet')

    def test_aggregate_sheets_and_blank_rows_are_not_records(self):
        def extras(w):
            w.create_sheet('Sheet1').append(['=1/0', 'PRIVATE SENTINEL'])
            w.create_sheet('Calc_Data')
            w.create_sheet('Statistics')
            w['Lattakia'].append([None] * 15)
        self.mutate(extras)
        self.assertEqual(self.run_import()[0]['diff']['after']['total'], 4)

    def test_formula_in_data_is_rejected_without_logging_it(self):
        self.mutate(lambda w: setattr(w['Homs']['O2'], 'value', '=SECRET("hidden")'))
        self.assert_blocked('formula_not_allowed')
        self.assertNotIn('SECRET', (self.review / 'report.json').read_text())

    def test_unused_serial_formula_is_ignored(self):
        self.mutate(lambda w: setattr(w['Homs']['A2'], 'value', '=ROW()-1'))
        self.assertEqual(self.run_import()[1], 0)

    def test_excel_error(self):
        self.mutate(lambda w: setattr(w['Homs']['M2'], 'value', '#VALUE!'))
        self.assert_blocked('excel_error')

    def test_partial_row_is_not_silently_skipped(self):
        self.mutate(lambda w: w['Homs'].append([None, 'PRIVATE SENTINEL']))
        self.assert_blocked('missing_facility_code')

    def test_unknown_program_and_governorate_mismatch(self):
        self.rows['Homs'][0]['main_program'] = 'Typo'
        self.rows['Homs'][0]['governorate'] = 'Hama'
        self.write_workbook()
        report, code = self.run_import()
        self.assertEqual(code, 2)
        self.assertIn('unknown_program', report['errors'])
        self.assertIn('governorate_sheet_mismatch', report['errors'])

    def test_duplicate_id_in_other_sheet_case_and_whitespace(self):
        self.rows['Homs'][0]['facility_code'] = '  test-0 '
        self.write_workbook()
        self.assert_blocked('duplicate_facility_code')

    def test_duplicate_content_with_different_id(self):
        twin = dict(self.rows['Homs'][0], facility_code='Another-code')
        self.rows['Homs'].append(twin)
        self.write_workbook()
        self.assert_blocked('duplicate_record_different_code')

    def test_colocated_distinct_facilities_stay_separate(self):
        twin = dict(self.rows['Homs'][0], facility_code='Another-code', implementing_partner='Other Partner')
        self.rows['Homs'].append(twin)
        policy = json.loads(self.policy.read_text())
        policy['approved_public_codes'].append('Another-code')
        self.policy.write_text(imp.dumps(policy))
        self.write_workbook()
        report, code = self.run_import()
        self.assertEqual(code, 0)
        self.assertEqual(report['diff']['after']['mapped'], 5)
        _, records = imp.read_baseline(self.review / 'facilities-data.js')
        colocated = [r for r in records if r['governorate_en'] == 'Homs']
        self.assertEqual(len({r['id'] for r in colocated}), 2)
        self.assertEqual(colocated[0]['lat'], colocated[1]['lat'])
        self.assertTrue(all(any('Shares identical' in f for f in r['quality_flags']) for r in colocated))

    def test_decimal_dms_and_invalid_coordinates(self):
        self.assertAlmostEqual(imp.coordinate('35°30\'30.0"N', 'lat'), 35.5083333333)
        self.assertAlmostEqual(imp.coordinate('36°30\'30.0', 'lon'), 36.5083333333)
        self.assertEqual(imp.coordinate('\u00a035.9', 'lon'), 35.9)
        for value, axis in [('NaN', 'lat'), ('inf', 'lon'), (True, 'lat'),
                            ('91', 'lat'), ('181', 'lon'), (None, 'lat'),
                            ('35°60\'0', 'lat'), ('35°1\'60', 'lat'),
                            ('35°1\'0E', 'lat'), ('-35°1\'0N', 'lat'),
                            ('35°1\'0 junk', 'lat'), ('35,8', 'lat')]:
            with self.subTest(value=value):
                with self.assertRaises(imp.Invalid):
                    imp.coordinate(value, axis)

    def test_missing_and_outside_coordinates_block_even_protected(self):
        self.rows['Homs'][0].update(main_program='GBV', lat=None)
        self.write_workbook()
        self.assert_blocked('missing_coordinate')
        self.rows['Homs'][0]['lat'] = 0
        self.write_workbook()
        self.assert_blocked('outside_syria_extent_review_source')

    def test_existing_swap_retained_only_when_values_match(self):
        self.old[0].update(lat=35.3, lon=35.9, lat_original='35.9', lon_original='35.3',
                           quality_flags=[imp.SWAP_FLAG])
        self.baseline.write_text(imp.serialize(self.old))
        self.rows['Lattakia'][0].update(lat=35.9, lon=35.3)
        self.write_workbook()
        report, code = self.run_import()
        self.assertEqual(code, 0)
        self.assertEqual(report['warnings']['existing_unverified_swap_retained'], 1)
        self.rows['Lattakia'][0]['lat'] = 35.8
        self.write_workbook()
        self.assert_blocked('outside_syria_extent_review_source')

    def test_confirmed_source_correction_clears_legacy_swap(self):
        self.old[0].update(lat_original='36.5', lon_original='35.2', quality_flags=[imp.SWAP_FLAG])
        self.baseline.write_text(imp.serialize(self.old))
        report, code = self.run_import()
        self.assertEqual(code, 0)
        self.assertNotIn('existing_unverified_swap_retained', report['warnings'])

    def test_protected_fields_cannot_leak_to_candidate_report_or_console(self):
        secret = 'UNIQUE-PRIVATE-SENTINEL'
        protected = dict(self.rows['Homs'][0])
        for k in ('facility_code', 'implementing_partner', 'district', 'subdistrict',
                  'commune', 'neighbourhood', 'site_name', 'brochure'):
            protected[k] = secret + k
        protected.update(main_program='GBV', facility_type='Women and Girls Safe Space',
                         mode='WGSS', lat=35.654321, lon=36.123456)
        self.rows['Homs'].append(protected)
        self.write_workbook()
        capture = io.StringIO()
        with redirect_stdout(capture):
            code = imp.main([str(self.xlsx), '--baseline', str(self.baseline),
                             '--policy', str(self.policy), '--review-dir', str(self.review)])
        self.assertEqual(code, 0)
        for blob in [capture.getvalue(), *(p.read_text() for p in self.review.iterdir())]:
            for value in [secret, '35.654321', '36.123456']:
                self.assertNotIn(value, blob)
        _, records = imp.read_baseline(self.review / 'facilities-data.js')
        r = next(r for r in records if r.get('protected'))
        self.assertEqual(r, imp.protected_record('Homs', 'GBV', 1))

    def test_rh_wgss_and_unknown_type_are_withheld_despite_approval(self):
        for kind, mode in [('Women and Girls Safe Space', 'WGSS'), ('Unknown facility', 'RH SC')]:
            self.rows['Homs'][0].update(facility_type=kind, mode=mode)
            self.write_workbook()
            report, code = self.run_import()
            self.assertEqual(code, 0)
            self.assertEqual(report['diff']['after']['protected'], 1)

    def test_unapproved_new_rh_is_withheld(self):
        self.rows['Homs'][0]['facility_code'] = 'NEW-PUBLIC-NOT-APPROVED'
        self.write_workbook()
        report, _ = self.run_import()
        self.assertEqual(report['diff']['after']['protected'], 1)
        self.assertNotIn('NEW-PUBLIC-NOT-APPROVED', (self.review / 'facilities-data.js').read_text())

    def test_protected_edits_and_row_order_do_not_change_public_bytes(self):
        self.rows['Homs'][0]['main_program'] = 'GBV'
        self.write_workbook()
        self.run_import()
        before = (self.review / 'facilities-data.js').read_bytes()
        self.rows['Homs'][0].update(facility_code='OTHER-SECRET', lat=35.45678, site_name='OTHER SECRET')
        self.write_workbook()
        report, _ = self.run_import()
        self.assertEqual(before, (self.review / 'facilities-data.js').read_bytes())
        self.assertNotIn('OTHER', imp.dumps(report))

    def test_count_changes_and_substitutions_are_reported(self):
        self.rows['Homs'][0]['facility_code'] = 'replacement'
        self.write_workbook()
        report, _ = self.run_import()
        self.assertEqual(report['diff']['before']['total'], report['diff']['after']['total'])
        self.assertEqual(report['diff']['public_records_removed_or_newly_withheld'], 1)
        self.rows['Homs'].append(dict(self.rows['Homs'][0], facility_code='additional', site_name='Other'))
        self.write_workbook()
        report, _ = self.run_import()
        self.assertEqual(report['diff']['after']['total'], 5)
        self.assertIn('facility_counts_changed_review_required', report['warnings'])

    def test_same_count_governorate_changes_are_reported(self):
        moved = self.rows['Homs'][0]
        self.rows['Homs'].append(dict(moved, facility_code='second', site_name='Other'))
        self.write_workbook()
        r, _ = self.run_import()
        self.run_import(apply=True, approve=r['approval_token'])
        self.rows['Homs'].pop()
        self.rows['Hama'].append(dict(moved, facility_code='second', governorate='Hama', site_name='Other'))
        self.write_workbook()
        r, _ = self.run_import()
        self.assertEqual(r['diff']['before']['total'], r['diff']['after']['total'])
        self.assertIn('facility_counts_changed_review_required', r['warnings'])

    def test_failed_validation_removes_stale_preview(self):
        self.run_import()
        self.rows['Homs'][0]['lat'] = 'bad'
        self.write_workbook()
        self.assert_blocked('invalid_coordinate')

    def test_arbitrary_javascript_baseline_is_never_executed(self):
        self.baseline.write_text('throw new Error("SECRET")')
        self.assert_blocked('invalid_public_baseline')

    def test_raw_nonpublic_baseline_rejected(self):
        self.old[0]['main_program'] = 'GBV'
        self.baseline.write_text(imp.serialize(self.old))
        self.assert_blocked('unsafe_public_baseline')

    def test_deterministic_candidate(self):
        self.run_import()
        one = (self.review / 'facilities-data.js').read_bytes()
        self.run_import()
        self.assertEqual(one, (self.review / 'facilities-data.js').read_bytes())

    def test_report_write_failure_cannot_replace_baseline(self):
        report, _ = self.run_import()
        before = self.baseline.read_bytes()
        with patch.object(imp, 'atomic_write', side_effect=OSError('PRIVATE SENTINEL')):
            with redirect_stdout(io.StringIO()) as capture:
                code = imp.main([str(self.xlsx), '--baseline', str(self.baseline),
                                 '--policy', str(self.policy), '--review-dir', str(self.review),
                                 '--apply', '--approve', report['approval_token']])
        self.assertEqual(code, 2)
        self.assertEqual(self.baseline.read_bytes(), before)
        self.assertNotIn('PRIVATE SENTINEL', capture.getvalue())

    def test_invalid_zip_has_sanitized_error(self):
        self.xlsx.write_text('PRIVATE SENTINEL')
        with redirect_stdout(io.StringIO()) as capture:
            code = imp.main([str(self.xlsx), '--baseline', str(self.baseline),
                             '--policy', str(self.policy), '--review-dir', str(self.review)])
        self.assertEqual(code, 2)
        self.assertNotIn('PRIVATE SENTINEL', capture.getvalue())
        self.assertIn('input_or_output_failure', capture.getvalue())

    def test_source_output_collision(self):
        with self.assertRaisesRegex(imp.Invalid, 'input_output_path_collision'):
            imp.run(self.xlsx, self.review / 'facilities-data.js', self.policy, self.review)


if __name__ == '__main__':
    unittest.main()
